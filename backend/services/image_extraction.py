"""Spreadsheet-driven public image downloader for the Image Extraction tool."""
import asyncio
import csv
import ipaddress
import io
import json
import os
import re
import shutil
import socket
import tempfile
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urljoin, urlparse

import httpx
from fastapi import APIRouter, File, Form, HTTPException, UploadFile, Request
from fastapi.responses import FileResponse, StreamingResponse
from starlette.background import BackgroundTask, BackgroundTasks
from . import auth_store
from .session_cleanup import operation_lease

router = APIRouter()
MAX_SHEET_BYTES = 20 * 1024 * 1024
MAX_IMAGE_BYTES = 15 * 1024 * 1024
MAX_REDIRECTS = 5
MAX_DOWNLOADS = 20000
ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
RESULTS_DIR = os.path.abspath(os.environ.get("EXTRACTION_TEMP_DIR", os.environ.get("SESSION_BASE_DIR", os.path.join(os.path.dirname(__file__), "..", "..", "output", "sessions"))))


def _read_sheet(filename: str, content: bytes):
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Upload an .xlsx or .csv file")
    if len(content) > MAX_SHEET_BYTES:
        raise HTTPException(413, "Spreadsheet exceeds the 20 MB limit")
    try:
        if ext == ".csv":
            text = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text, newline="")))
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
    except Exception as exc:
        raise HTTPException(400, f"Could not read spreadsheet: {exc}")
    if not rows:
        raise HTTPException(400, "The spreadsheet is empty")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    columns = [{"index": i, "name": header or f"Column {i + 1}"} for i, header in enumerate(headers)]
    return rows, columns


@router.post("/api/extract/columns")
async def get_columns(file: UploadFile = File(...)):
    content = await file.read(MAX_SHEET_BYTES + 1)
    _, columns = _read_sheet(file.filename or "", content)
    return {"columns": columns}


def _validate_public_url(url: str):
    parsed = urlparse(url)
    if parsed.scheme.lower() not in ("http", "https") or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Only public http(s) image URLs are allowed")
    try:
        addresses = {result[4][0] for result in socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)}
    except OSError:
        raise ValueError("Could not resolve host")
    if not addresses:
        raise ValueError("Could not resolve host")
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%", 1)[0])
        if not ip.is_global:
            raise ValueError("Private or reserved network addresses are not allowed")
    return parsed


def _download_one(client: httpx.Client, url: str, output_path: str):
    current = url.strip()
    try:
        for redirect in range(MAX_REDIRECTS + 1):
            _validate_public_url(current)
            with client.stream("GET", current, follow_redirects=False) as response:
                if response.is_redirect:
                    if redirect >= MAX_REDIRECTS:
                        raise ValueError("Too many redirects")
                    target = response.headers.get("location")
                    if not target:
                        raise ValueError("Redirect has no destination")
                    current = urljoin(current, target)
                    continue
                response.raise_for_status()
                mime = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
                ext_by_type = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
                url_ext = os.path.splitext(urlparse(current).path)[1].lower()
                ext = ext_by_type.get(mime, url_ext if url_ext in (".jpg", ".jpeg", ".png", ".webp") else ".jpg")
                size = 0
                with open(output_path, "wb") as out:
                    for chunk in response.iter_bytes():
                        size += len(chunk)
                        if size > MAX_IMAGE_BYTES:
                            raise ValueError("Image exceeds the 15 MB limit")
                        out.write(chunk)
                if size == 0:
                    raise ValueError("Image response was empty")
                with open(output_path, "rb") as image_file:
                    signature = image_file.read(12)
                valid = (signature.startswith(b"\xff\xd8\xff") if ext in (".jpg", ".jpeg") else
                         signature.startswith(b"\x89PNG\r\n\x1a\n") if ext == ".png" else
                         signature.startswith(b"RIFF") and signature[8:12] == b"WEBP" if ext == ".webp" else False)
                if not valid:
                    raise ValueError("Downloaded file is not a valid JPEG, PNG, or WebP image")
                return ext, None
        raise ValueError("Too many redirects")
    except Exception as exc:
        try:
            os.remove(output_path)
        except OSError:
            pass
        return None, str(exc)


def _safe_sheet_stem(filename: str):
    stem = os.path.splitext(os.path.basename(filename))[0]
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-")
    return safe or "sheet"


@router.post("/api/extract/run")
async def run_extraction(request: Request, file: UploadFile = File(...), column_index: int = Form(...)):
    content = await file.read(MAX_SHEET_BYTES + 1)
    rows, columns = _read_sheet(file.filename or "", content)
    if column_index < 0 or column_index >= len(columns):
        raise HTTPException(400, "Choose a valid column")
    os.makedirs(RESULTS_DIR, exist_ok=True)
    job_id = str(uuid.uuid4())
    auth_store.bind_extraction_job(job_id, request.state.user["email"])
    job_dir = tempfile.mkdtemp(prefix=f"extract_{job_id}_", dir=RESULTS_DIR)
    zip_path = os.path.join(job_dir, f"{_safe_sheet_stem(file.filename or 'sheet')}_images.zip")

    links = []
    links_per_row = {}
    failed = []
    for excel_row, row in enumerate(rows[1:], start=2):
        value = row[column_index] if column_index < len(row) else None
        if value is None or not str(value).strip():
            continue
        for raw_url in str(value).split(","):
            url = raw_url.strip()
            if not url:
                continue
            link_number = links_per_row.get(excel_row, 0) + 1
            links_per_row[excel_row] = link_number
            if not url.lower().startswith(("http://", "https://")):
                failed.append({"row": excel_row, "link": link_number, "url": url, "error": "Not a valid http(s) URL"})
            else:
                links.append((excel_row, link_number, url))
    if len(links) > MAX_DOWNLOADS:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(413, f"The spreadsheet contains more than {MAX_DOWNLOADS} image links")

    async def stream_events():
        completed = 0
        saved = 0
        total = len(links)
        client = httpx.Client(timeout=httpx.Timeout(30.0, connect=10.0), headers={"User-Agent": "NeosisImageExtractor/1.0"}, trust_env=False)
        downloads_dir = os.path.join(job_dir, "downloaded")
        os.makedirs(downloads_dir, exist_ok=True)
        try:
            lease = operation_lease(job_id, "extraction_job", "generation")
            lease.__enter__()
            with ThreadPoolExecutor(max_workers=20) as pool:
                loop = asyncio.get_running_loop()
                futures = [loop.run_in_executor(pool, lambda item=item: _download_with_client(client, item, downloads_dir)) for item in links]
                with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as archive:
                    for future in asyncio.as_completed(futures):
                        row_num, link_num, url, local_path, ext, error = await future
                        completed += 1
                        if error:
                            failed.append({"row": row_num, "link": link_num, "url": url, "error": error})
                        else:
                            base = f"image_{row_num}" if links_per_row[row_num] == 1 else f"image_{row_num}_{link_num}"
                            archive.write(local_path, f"{base}{ext}")
                            os.remove(local_path)
                            saved += 1
                        yield f"data: {json.dumps({'progress': completed, 'total': total})}\n\n"
            yield f"data: {json.dumps({'done': True, 'saved': saved, 'failed': failed, 'download_url': f'/api/extract/download/{job_id}', 'filename': os.path.basename(zip_path)})}\n\n"
        except Exception as exc:
            shutil.rmtree(job_dir, ignore_errors=True)
            yield f"data: {json.dumps({'error': str(exc)})}\n\n"
        finally:
            client.close()
            if 'lease' in locals():
                lease.__exit__(None, None, None)

    return StreamingResponse(stream_events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


def _download_with_client(client, item, downloads_dir):
    row_num, link_num, url = item
    local_path = os.path.join(downloads_dir, uuid.uuid4().hex)
    ext, error = _download_one(client, url, local_path)
    return row_num, link_num, url, local_path, ext, error


@router.get("/api/extract/download/{job_id}")
async def download_extraction(job_id: str, request: Request):
    if not re.fullmatch(r"[0-9a-f-]{36}", job_id):
        raise HTTPException(404, "Download not found")
    if not auth_store.owns_extraction_job(job_id, request.state.user["email"]):
        raise HTTPException(404, "Download not found")
    auth_store.touch_extraction_job(job_id)
    matches = [name for name in os.listdir(RESULTS_DIR) if name.startswith(f"extract_{job_id}_")]
    if not matches:
        raise HTTPException(404, "Download not found or expired")
    job_dir = os.path.join(RESULTS_DIR, matches[0])
    zips = [name for name in os.listdir(job_dir) if name.endswith("_images.zip")]
    if not zips:
        raise HTTPException(404, "Download not ready")
    path = os.path.join(job_dir, zips[0])
    lease = operation_lease(job_id, "extraction_job", "download")
    lease.__enter__()
    tasks = BackgroundTasks()
    tasks.add_task(lease.__exit__, None, None, None)
    tasks.add_task(_cleanup_extraction, job_dir, job_id)
    return FileResponse(path, media_type="application/zip", filename=zips[0], background=tasks)


def _cleanup_extraction(job_dir, job_id):
    shutil.rmtree(job_dir, ignore_errors=True)
    auth_store.delete_extraction_job(job_id)
