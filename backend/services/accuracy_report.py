"""Temporary, session-owned Accuracy Report preparation and Google Drive delivery."""
import asyncio
import io
import json
import logging
import os
import re
import shutil
import tempfile
import uuid
import warnings
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import quote, urlsplit

import httpx
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, StreamingResponse
from PIL import Image, ImageOps

from services import auth_store, auth
from services.session_manager import create_session_id, get_session_dir, validate_session
from services.session_cleanup import operation_lease

router = APIRouter(prefix="/api/accuracy-report", tags=["Accuracy Report"])
MAX_WORKBOOK_BYTES = 20 * 1024 * 1024
ALLOWED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.file"
GOOGLE_DRIVE_RESUMABLE = "https://www.googleapis.com/upload/drive/v3/files?uploadType=resumable&ignoreDefaultVisibility=true&fields=id,webViewLink,mimeType"
GOOGLE_DRIVE_DOWNLOAD = "https://www.googleapis.com/drive/v3/files/{file_id}/download?mimeType=application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GOOGLE_DRIVE_OPERATIONS = "https://www.googleapis.com/drive/v3/operations/{}"
logger = logging.getLogger(__name__)


def _limit(name, default, floor, ceiling):
    try:
        return max(floor, min(ceiling, int(os.environ.get(name, default))))
    except (TypeError, ValueError):
        return int(default)


def limits():
    return {"max_images": _limit("ACCURACY_REPORT_MAX_IMAGES", 200, 1, 200),
            "max_image_bytes": _limit("ACCURACY_REPORT_MAX_IMAGE_BYTES", 15 * 1024 * 1024, 1024, 100 * 1024 * 1024),
            "max_upload_bytes": _limit("ACCURACY_REPORT_MAX_UPLOAD_BYTES", 512 * 1024 * 1024, 1024 * 1024, 2 * 1024 * 1024 * 1024),
            "max_dimension": _limit("ACCURACY_REPORT_MAX_IMAGE_DIMENSION", 1600, 256, 4096),
            "jpeg_quality": _limit("ACCURACY_REPORT_JPEG_QUALITY", 82, 50, 95)}


def _owned(session_id, request):
    if not validate_session(session_id, request.state.user["email"]):
        raise HTTPException(404, "Report session not found")
    return Path(get_session_dir(session_id))


def _natural_key(value):
    return [(1, int(part)) if part.isdigit() else (0, part.casefold()) for part in re.split(r"(\d+)", value)]


def _image_is_readable(content):
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                image.verify()
        return True
    except Exception:
        return False


def _workbook_info(content):
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        result = {}
        for sheet in workbook.worksheets:
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            result[sheet.title] = [{"index": index, "name": str(value).strip() if value is not None else ""}
                                   for index, value in enumerate(first)]
        workbook.close()
        if not result:
            raise ValueError()
        return result
    except Exception:
        raise HTTPException(400, "Could not read the Excel workbook")


def _read_rows(content, sheet_name, column):
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise HTTPException(400, "Select a worksheet from the uploaded workbook")
        sheet = workbook[sheet_name]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if column < 0 or column >= len(headers):
            workbook.close()
            raise HTTPException(400, "Select a prediction column from the worksheet")
        rows = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = tuple(values)
            if any(value is not None and str(value).strip() for value in values):
                rows.append({"row_number": row_number, "prediction": values[column] if column < len(values) else None})
        workbook.close()
        return rows
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Could not read the selected worksheet")


def _parse_prediction(value):
    if not isinstance(value, str):
        return None
    try:
        parsed = json.loads(value)
    except (ValueError, TypeError):
        return None
    if not isinstance(parsed, list):
        return None
    totals = {}
    for item in parsed:
        if not isinstance(item, dict):
            return None
        name = item.get("Product")
        if name is None or not str(name).strip():
            name = item.get("Label")
        if name is None or not str(name).strip():
            return None
        count = item.get("Count")
        if isinstance(count, bool) or not isinstance(count, (int, float)) or int(count) != count or count < 0:
            return None
        sku_name = str(name).strip()
        totals[sku_name] = totals.get(sku_name, 0) + int(count)
    return [{"sku_name": name, "detected": detected, "available": detected} for name, detected in totals.items()]


def build_format2_preview(rows, image_records):
    by_row = {}
    duplicate_names = []
    for image in image_records:
        match = re.fullmatch(r"image_(\d+)", Path(image["name"]).stem, re.IGNORECASE)
        if not match:
            continue
        row_number = int(match.group(1))
        by_row.setdefault(row_number, []).append(image)
    duplicate_rows = {row: entries for row, entries in by_row.items() if len(entries) > 1}
    duplicate_names = [entry["name"] for entries in duplicate_rows.values() for entry in entries]
    row_map = {row["row_number"]: row for row in rows}
    row_without_image = [row for row in rows if row["row_number"] not in by_row or row["row_number"] in duplicate_rows]
    matched_image_keys = {entries[0]["key"] for number, entries in by_row.items()
                          if number in row_map and len(entries) == 1}
    image_without_row = [image["name"] for image in image_records if image["key"] not in matched_image_keys]
    blocks = []
    for row in rows:
        images = by_row.get(row["row_number"], [])
        if len(images) != 1:
            continue
        block = {"row_number": row["row_number"], "image_key": images[0]["key"], "image_name": images[0]["name"],
                 "items": None, "invalid_prediction": False}
        block["items"] = _parse_prediction(row["prediction"])
        block["invalid_prediction"] = block["items"] is None
        if block["items"] is None:
            block["items"] = []
        block["total_available"] = sum(item["available"] for item in block["items"])
        block["total_detected"] = sum(item["detected"] for item in block["items"])
        blocks.append(block)
    blocks.sort(key=lambda block: block["row_number"])
    return {"blocks": blocks, "skipped_rows": [row["row_number"] for row in row_without_image],
            "skipped_images": sorted(set(image_without_row), key=_natural_key), "duplicate_images": sorted(duplicate_names, key=_natural_key),
            "matched_count": len(blocks), "row_count": len(rows), "image_count": len(image_records),
            "invalid_rows": [block["row_number"] for block in blocks if block["invalid_prediction"]]}


def build_format1_preview(image_records):
    good = sorted(image_records, key=lambda entry: _natural_key(entry["name"]))
    return {"blocks": [{"row_number": None, "image_key": entry["key"], "image_name": entry["name"], "items": [],
                        "invalid_prediction": False, "total_available": None, "total_detected": None}
                       for entry in good],
            "skipped_images": [], "unreadable_images": [], "matched_count": len(good), "row_count": 0,
            "image_count": len(good), "invalid_rows": [], "skipped_rows": [], "duplicate_images": []}


def _processed_image(source, destination, config):
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(source) as opened:
            image = ImageOps.exif_transpose(opened)
            image.thumbnail((config["max_dimension"], config["max_dimension"]), Image.Resampling.LANCZOS)
            if image.mode not in {"RGB", "L"}:
                rgba = image.convert("RGBA")
                background = Image.new("RGB", rgba.size, "white")
                background.paste(rgba, mask=rgba.getchannel("A"))
                image = background
            elif image.mode == "L":
                image = image.convert("RGB")
            image.save(destination, format="JPEG", quality=config["jpeg_quality"], optimize=True)


def _build_docx(preview, image_dir, title, output_file, image_urls, config):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
          "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
          "ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
    for prefix, uri in ns.items():
        if prefix not in {"rel", "ct"}:
            ET.register_namespace(prefix, uri)
    q = lambda prefix, name: "{" + ns[prefix] + "}" + name
    doc = ET.Element(q("w", "document"))
    body = ET.SubElement(doc, q("w", "body"))
    image_rels = []
    media = []

    def paragraph(text="", *, bold=False, size=None, page_break=False):
        p = ET.SubElement(body, q("w", "p"))
        p_props = ET.SubElement(p, q("w", "pPr"))
        ET.SubElement(p_props, q("w", "jc"), {q("w", "val"): "left"})
        if page_break:
            run = ET.SubElement(p, q("w", "r"))
            ET.SubElement(run, q("w", "br"), {q("w", "type"): "page"})
        if text:
            run = ET.SubElement(p, q("w", "r"))
            props = ET.SubElement(run, q("w", "rPr"))
            if bold:
                ET.SubElement(props, q("w", "b"))
            if size:
                ET.SubElement(props, q("w", "sz"), {q("w", "val"): str(size)})
            ET.SubElement(run, q("w", "t")).text = text
        return p

    def cell_text(cell, text, bold=False):
        p = ET.SubElement(cell, q("w", "p"))
        p_props = ET.SubElement(p, q("w", "pPr"))
        ET.SubElement(p_props, q("w", "jc"), {q("w", "val"): "left"})
        r = ET.SubElement(p, q("w", "r"))
        if bold:
            ET.SubElement(ET.SubElement(r, q("w", "rPr")), q("w", "b"))
        ET.SubElement(r, q("w", "t")).text = text

    def add_picture(path, index):
        rel_id = "rId" + str(index + 1)
        with Image.open(path) as image:
            width, height = image.size
        # Letter page with 0.5in margins: keep each image within a generous
        # content box so the image itself cannot spill onto a second page.
        max_width_emu, max_height_emu = 6858000, 5486400
        scale = min(max_width_emu / width, max_height_emu / height)
        width_emu = max(1, int(width * scale))
        height_emu = max(1, int(height * scale))
        image_paragraph = ET.SubElement(body, q("w", "p"))
        p_props = ET.SubElement(image_paragraph, q("w", "pPr"))
        ET.SubElement(p_props, q("w", "jc"), {q("w", "val"): "left"})
        drawing = ET.SubElement(ET.SubElement(image_paragraph, q("w", "r")), q("w", "drawing"))
        inline = ET.SubElement(drawing, q("wp", "inline"), {"distT": "0", "distB": "0", "distL": "0", "distR": "0"})
        ET.SubElement(inline, q("wp", "extent"), {"cx": str(width_emu), "cy": str(height_emu)})
        ET.SubElement(inline, q("wp", "docPr"), {"id": str(index + 1), "name": f"Image {index + 1}"})
        graphic = ET.SubElement(inline, q("a", "graphic"))
        graphic_data = ET.SubElement(graphic, q("a", "graphicData"), {"uri": ns["pic"]})
        pic = ET.SubElement(graphic_data, q("pic", "pic"))
        nv = ET.SubElement(pic, q("pic", "nvPicPr"))
        ET.SubElement(nv, q("pic", "cNvPr"), {"id": "0", "name": f"Image {index + 1}.jpg"})
        ET.SubElement(nv, q("pic", "cNvPicPr"))
        fill = ET.SubElement(pic, q("pic", "blipFill"))
        ET.SubElement(fill, q("a", "blip"), {q("r", "embed"): rel_id})
        stretch = ET.SubElement(fill, q("a", "stretch")); ET.SubElement(stretch, q("a", "fillRect"))
        shape = ET.SubElement(pic, q("pic", "spPr"))
        transform = ET.SubElement(shape, q("a", "xfrm")); ET.SubElement(transform, q("a", "off"), {"x": "0", "y": "0"}); ET.SubElement(transform, q("a", "ext"), {"cx": str(width_emu), "cy": str(height_emu)})
        geometry = ET.SubElement(shape, q("a", "prstGeom"), {"prst": "rect"}); ET.SubElement(geometry, q("a", "avLst"))
        image_rels.append((rel_id, f"media/image{index + 1}.jpg"))
        media.append((f"word/media/image{index + 1}.jpg", Path(path).read_bytes()))

    def add_table(rows):
        table = ET.SubElement(body, q("w", "tbl"))
        props = ET.SubElement(table, q("w", "tblPr"))
        borders = ET.SubElement(props, q("w", "tblBorders"))
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            ET.SubElement(borders, q("w", edge), {q("w", "val"): "single", q("w", "sz"): "4", q("w", "space"): "0", q("w", "color"): "999999"})
        grid = ET.SubElement(table, q("w", "tblGrid"))
        for _ in range(4): ET.SubElement(grid, q("w", "gridCol"), {q("w", "w"): "2200"})
        for row_index, values in enumerate(rows):
            tr = ET.SubElement(table, q("w", "tr"))
            for value in values:
                tc = ET.SubElement(tr, q("w", "tc")); cell_text(tc, str(value), bold=row_index == 0 or value == "Total")

    paragraph(title, bold=True, size=36)
    for index, block in enumerate(preview["blocks"], start=1):
        paragraph(f"Image {index}", bold=True, size=30, page_break=index > 1)
        add_picture(image_urls[block["image_key"]], index - 1)
        valid_format2 = preview["format"] == 2 and not block["invalid_prediction"]
        values = [("Total SKU Available", str(block["total_available"]) if valid_format2 else ""),
                  ("SKU Detected", str(block["total_detected"]) if valid_format2 else ""),
                  ("Incorrect", ""), ("Missing", ""), ("Accuracy", "")]
        for label, value in values:
            p = ET.SubElement(body, q("w", "p")); p_props = ET.SubElement(p, q("w", "pPr")); ET.SubElement(p_props, q("w", "jc"), {q("w", "val"): "left"}); r = ET.SubElement(p, q("w", "r")); ET.SubElement(ET.SubElement(r, q("w", "rPr")), q("w", "b")); ET.SubElement(r, q("w", "t")).text = label + ": "
            if value:
                ET.SubElement(ET.SubElement(p, q("w", "r")), q("w", "t")).text = value
        if preview["format"] == 2:
            if block["invalid_prediction"]:
                paragraph("invalid prediction data")
            else:
                rows = [["SKU Name", "Detected", "Available", "Accuracy"]]
                rows.extend([[item["sku_name"], item["detected"], item["available"], ""] for item in block["items"]])
                rows.append(["Total", block["total_detected"], block["total_available"], ""])
                add_table(rows)
    section = ET.SubElement(body, q("w", "sectPr"))
    ET.SubElement(section, q("w", "pgSz"), {q("w", "w"): "12240", q("w", "h"): "15840"})
    ET.SubElement(section, q("w", "pgMar"), {q("w", "top"): "720", q("w", "right"): "720",
                                                 q("w", "bottom"): "720", q("w", "left"): "720",
                                                 q("w", "header"): "360", q("w", "footer"): "360",
                                                 q("w", "gutter"): "0"})
    document_xml = ET.tostring(doc, encoding="utf-8", xml_declaration=True)
    content_types = ET.Element(q("ct", "Types"))
    ET.SubElement(content_types, q("ct", "Default"), {"Extension": "rels", "ContentType": "application/vnd.openxmlformats-package.relationships+xml"})
    ET.SubElement(content_types, q("ct", "Default"), {"Extension": "xml", "ContentType": "application/xml"})
    ET.SubElement(content_types, q("ct", "Default"), {"Extension": "jpg", "ContentType": "image/jpeg"})
    ET.SubElement(content_types, q("ct", "Override"), {"PartName": "/word/document.xml", "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"})
    root_rels = ET.Element(q("rel", "Relationships"))
    ET.SubElement(root_rels, q("rel", "Relationship"), {"Id": "rId1", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", "Target": "word/document.xml"})
    doc_rels = ET.Element(q("rel", "Relationships"))
    for rel_id, target in image_rels:
        ET.SubElement(doc_rels, q("rel", "Relationship"), {"Id": rel_id, "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image", "Target": target})
    with zipfile.ZipFile(output_file, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", ET.tostring(content_types, encoding="utf-8", xml_declaration=True))
        archive.writestr("_rels/.rels", ET.tostring(root_rels, encoding="utf-8", xml_declaration=True))
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", ET.tostring(doc_rels, encoding="utf-8", xml_declaration=True))
        for name, content in media:
            archive.writestr(name, content)


async def _read_upload(upload, max_bytes):
    content = await upload.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise HTTPException(413, "Uploaded file exceeds the configured size limit")
    return content


@router.post("/sessions")
async def create_report_session(request: Request):
    session_id = create_session_id()
    auth_store.bind_tool_session(session_id, request.state.user["email"], "accuracy_report")
    directory = Path(get_session_dir(session_id))
    (directory / "accuracy_report").mkdir(parents=True, exist_ok=True)
    (directory / "accuracy_report" / "cleanup-tag.json").write_text(json.dumps({"session_type": "accuracy_report", "temporary": True}), encoding="utf-8")
    return {"session_id": session_id, "limits": limits()}


@router.post("/sessions/{session_id}/workbook")
async def save_workbook(session_id: str, request: Request, file: UploadFile = File(...)):
    directory = _owned(session_id, request) / "accuracy_report"
    if Path(file.filename or "").suffix.casefold() != ".xlsx":
        raise HTTPException(400, "Upload an .xlsx workbook")
    content = await _read_upload(file, MAX_WORKBOOK_BYTES)
    sheets = _workbook_info(content)
    (directory / "workbook.xlsx").write_bytes(content)
    return {"worksheets": [{"name": name, "columns": columns} for name, columns in sheets.items()]}


@router.post("/sessions/{session_id}/images")
async def save_images(session_id: str, request: Request, files: list[UploadFile] = File(...)):
    directory = _owned(session_id, request) / "accuracy_report"
    config = limits()
    if len(files) > config["max_images"]:
        raise HTTPException(413, f"A report can include up to {config['max_images']} images")
    images_dir = directory / "images"
    if images_dir.exists():
        shutil.rmtree(images_dir)
    images_dir.mkdir(parents=True)
    records, skipped, total_bytes = [], [], 0
    for index, upload in enumerate(files):
        name = (upload.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
        ext = Path(name).suffix.casefold()
        content = await _read_upload(upload, config["max_image_bytes"])
        total_bytes += len(content)
        if total_bytes > config["max_upload_bytes"]:
            raise HTTPException(413, "Image uploads exceed the configured report size limit")
        if ext not in ALLOWED_IMAGE_EXTENSIONS:
            skipped.append({"name": name, "reason": "unsupported file type"})
            continue
        if not _image_is_readable(content):
            skipped.append({"name": name, "reason": "unreadable image"})
            continue
        key = f"{index:04d}-{uuid.uuid4().hex}{ext}"
        (images_dir / key).write_bytes(content)
        records.append({"key": key, "name": name})
    (directory / "images.json").write_text(json.dumps({"records": records, "skipped": skipped}, ensure_ascii=False), encoding="utf-8")
    return {"images": len(records), "skipped": skipped, "limits": config}


@router.get("/sessions/{session_id}/worksheets")
async def workbook_worksheets(session_id: str, request: Request):
    directory = _owned(session_id, request) / "accuracy_report"
    workbook = directory / "workbook.xlsx"
    if not workbook.is_file():
        raise HTTPException(404, "Upload a workbook first")
    sheets = _workbook_info(workbook.read_bytes())
    return {"worksheets": [{"name": name, "columns": columns} for name, columns in sheets.items()]}


@router.post("/sessions/{session_id}/validate")
async def validate_report_inputs(session_id: str, request: Request, format: int = Form(...), worksheet: str = Form(""), column: int = Form(0)):
    directory = _owned(session_id, request) / "accuracy_report"
    state_file, images_file = directory / "images.json", directory / "workbook.xlsx"
    if format not in (1, 2):
        raise HTTPException(400, "Select a report format")
    if not state_file.is_file():
        raise HTTPException(400, "Upload an image directory first")
    saved = json.loads(state_file.read_text(encoding="utf-8"))
    records = saved["records"]
    if format == 1:
        preview = build_format1_preview(records)
        preview["format"] = 1
    else:
        if not images_file.is_file():
            raise HTTPException(400, "Upload an .xlsx workbook for Format 2")
        rows = _read_rows(images_file.read_bytes(), worksheet, column)
        preview = build_format2_preview(rows, records)
        preview["format"] = 2
        if preview["matched_count"] == 0:
            preview["generation_blocked"] = True
    preview["unreadable_images"] = saved.get("skipped", [])
    (directory / "preview.json").write_text(json.dumps(preview, ensure_ascii=False), encoding="utf-8")
    return preview


@router.get("/sessions/{session_id}/images/{image_key}")
async def report_image(session_id: str, image_key: str, request: Request):
    directory = _owned(session_id, request) / "accuracy_report" / "images"
    if Path(image_key).name != image_key:
        raise HTTPException(404, "Image not found")
    path = directory / image_key
    if not path.is_file():
        raise HTTPException(404, "Image not found")
    return FileResponse(path, headers={"Cache-Control": "private, no-store"})


@router.get("/google/authorize")
async def authorize_report_google(request: Request):
    params = request.query_params
    session_id = params.get("session_id", "")
    _owned(session_id, request)
    if not os.environ.get("GOOGLE_CLIENT_ID") or not os.environ.get("GOOGLE_CLIENT_SECRET") or not os.environ.get("OAUTH_REDIRECT_URI"):
        raise HTTPException(503, "Google document creation is not configured")
    state, nonce = auth.secrets.token_urlsafe(32), auth.secrets.token_urlsafe(32)
    verifier, challenge = auth._make_pkce_pair()
    auth_store.save_oauth_transaction(state, nonce, verifier, "__accuracy_report_oauth__")
    from urllib.parse import urlencode
    query = urlencode({"client_id": os.environ["GOOGLE_CLIENT_ID"], "redirect_uri": os.environ["OAUTH_REDIRECT_URI"],
                       "response_type": "code", "scope": "openid email profile " + DRIVE_SCOPE,
                       "state": state, "nonce": nonce, "code_challenge": challenge, "code_challenge_method": "S256",
                       "access_type": "online", "include_granted_scopes": "true", "prompt": "select_account"})
    return RedirectResponse(auth.GOOGLE_AUTH_URL + "?" + query, status_code=302)


def _oauth_popup_html(message):
    encoded = json.dumps(message).replace("<", "\\u003c")
    return f"<!doctype html><meta charset=utf-8><title>Google permission</title><p>You can close this window.</p><script>if(window.opener){{window.opener.postMessage({encoded},window.location.origin);window.close()}}</script>"


async def _drive_create(access_token, content, title):
    metadata = {"name": title, "mimeType": "application/vnd.google-apps.document"}
    async with httpx.AsyncClient(timeout=90, follow_redirects=False) as client:
        response = await client.post(GOOGLE_DRIVE_RESUMABLE, json=metadata,
            headers={"Authorization": "Bearer " + access_token, "Content-Type": "application/json; charset=UTF-8",
                     "X-Upload-Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     "X-Upload-Content-Length": str(len(content))})
        if response.status_code >= 400:
            raise RuntimeError(f"Drive resumable session returned HTTP {response.status_code}")
        upload_url = response.headers.get("Location")
        if not upload_url:
            raise RuntimeError("Google Drive did not start a resumable upload")
        parsed = urlsplit(upload_url)
        if parsed.scheme != "https" or not parsed.hostname or not (parsed.hostname == "www.googleapis.com" or parsed.hostname.endswith(".googleapis.com")):
            raise RuntimeError("Google Drive returned an invalid upload location")
        response = await client.put(upload_url, content=content,
            headers={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     "Content-Length": str(len(content))})
    if response.status_code >= 400:
        raise RuntimeError(f"Drive media upload returned HTTP {response.status_code}")
    data = response.json()
    file_id = data.get("id")
    if not file_id:
        raise RuntimeError("Google Drive did not return a document")
    return file_id, data.get("webViewLink") or "https://docs.google.com/document/d/" + file_id + "/edit"


def _safe_generation_failure(stage, exc):
    """Return an actionable, secret-free diagnostic for the UI and logs."""
    message = str(exc)
    if message == "Prepared document exceeds the Google Docs import size limit":
        return "document_too_large", "The prepared document is too large for Google Docs. Reduce the image count or image sizes and retry."
    match = re.fullmatch(r"Drive (resumable session|media upload) returned HTTP (\d{3})", message)
    if match:
        operation, status = match.groups()
        code = f"drive_{operation.replace(' ', '_')}_http_{status}"
        if status == "401":
            detail = "Google rejected the access token. Retry and approve Drive access with the same Google account."
        elif status == "403":
            detail = "Google denied document creation. Check that the Drive API is enabled and Drive access was granted."
        elif status == "404":
            detail = "Google Drive did not find the upload endpoint. Check the Drive API configuration."
        elif status == "429":
            detail = "Google Drive is rate limiting document creation. Wait briefly, then retry."
        elif int(status) >= 500:
            detail = "Google Drive is temporarily unavailable. Retry in a few minutes."
        else:
            detail = f"Google Drive rejected the {operation} (HTTP {status}). Check the Google Drive configuration and retry."
        return code, detail
    if isinstance(exc, httpx.TimeoutException):
        return "drive_timeout", "The document upload timed out while contacting Google Drive. Retry with fewer or smaller images."
    if isinstance(exc, httpx.TransportError):
        return "drive_unreachable", "The server could not reach Google Drive. Check outbound network access and retry."
    return "generation_failed", f"{stage} failed. Review the inputs and retry."


def _sse(event):
    return "data: " + json.dumps(event, ensure_ascii=False) + "\n\n"


@router.post("/sessions/{session_id}/generate")
async def generate_report(session_id: str, request: Request, title: str = Form("")):
    directory = _owned(session_id, request) / "accuracy_report"
    access_token = request.headers.get("authorization", "")
    if not access_token.startswith("Bearer ") or len(access_token) > 16384:
        raise HTTPException(401, "Google permission is required to create the document")
    access_token = access_token[7:]
    preview_file = directory / "preview.json"
    if not preview_file.is_file():
        raise HTTPException(400, "Validate report inputs before generating")
    preview = json.loads(preview_file.read_text(encoding="utf-8"))
    if not preview.get("blocks") or preview.get("generation_blocked"):
        raise HTTPException(400, "At least one matched image is required")
    final_title = (title.strip() or "Accuracy Report")[:180]

    async def events():
        temp_dir = Path(tempfile.mkdtemp(prefix="accuracy-report-"))
        complete = False
        current_stage = "Reading data"
        try:
            yield _sse({"stage": "Reading data", "state": "started"})
            preview["format"] = int(preview["format"])
            yield _sse({"stage": "Reading data", "state": "complete"})
            current_stage = "Mapping images"
            yield _sse({"stage": current_stage, "state": "started"})
            report_images = directory / "images"
            image_sources = {}
            for block in preview["blocks"]:
                source = report_images / block["image_key"]
                if not source.is_file():
                    raise RuntimeError("An uploaded image is no longer available")
                image_sources[block["image_key"]] = source
            yield _sse({"stage": current_stage, "state": "complete"})
            current_stage = "Preparing images"
            yield _sse({"stage": current_stage, "state": "started"})
            image_urls = {}
            for index, block in enumerate(preview["blocks"], start=1):
                destination = temp_dir / f"{index:04d}.jpg"
                await asyncio.to_thread(_processed_image, image_sources[block["image_key"]], destination, limits())
                image_urls[block["image_key"]] = str(destination)
                yield _sse({"stage": current_stage, "state": "progress", "completed": index, "total": len(preview["blocks"])})
            yield _sse({"stage": current_stage, "state": "complete"})
            current_stage = "Creating document"
            yield _sse({"stage": current_stage, "state": "started"})
            docx_path = temp_dir / "report.docx"
            await asyncio.to_thread(_build_docx, preview, report_images, final_title, docx_path, image_urls, limits())
            max_import_bytes = 45 * 1024 * 1024
            compression_profiles = [(min(limits()["max_dimension"], 1200), min(limits()["jpeg_quality"], 70)),
                                    (900, 60), (700, 50)]
            for dimension, quality in compression_profiles:
                if docx_path.stat().st_size <= max_import_bytes:
                    break
                for index, block in enumerate(preview["blocks"], start=1):
                    compressed = {"max_dimension": dimension, "jpeg_quality": quality}
                    await asyncio.to_thread(_processed_image, image_sources[block["image_key"]], temp_dir / f"{index:04d}.jpg", compressed)
                await asyncio.to_thread(_build_docx, preview, report_images, final_title, docx_path, image_urls, limits())
            if docx_path.stat().st_size > max_import_bytes:
                raise RuntimeError("Prepared document exceeds the Google Docs import size limit")
            doc_id, doc_url = await _drive_create(access_token, docx_path.read_bytes(), final_title)
            (directory / "result.json").write_text(json.dumps({"doc_id": doc_id, "doc_url": doc_url, "title": final_title}), encoding="utf-8")
            complete = True
            yield _sse({"stage": current_stage, "state": "complete"})
            yield _sse({"state": "complete", "doc_url": doc_url,
                        "download_url": f"/api/accuracy-report/sessions/{session_id}/download", "title": final_title})
        except Exception as exc:
            stage = current_stage
            message = str(exc)
            if message in {"At least one matched image is required", "Report data is not available"}:
                stage = "Reading data"
            elif message == "An uploaded image is no longer available":
                stage = "Mapping images"
            failure_code, safe_message = _safe_generation_failure(stage, exc)
            logger.warning("Accuracy report generation failed at %s (%s; %s)", stage, type(exc).__name__, failure_code)
            yield _sse({"state": "error", "stage": stage, "message": safe_message})
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)
            if not complete:
                (directory / "result.json").unlink(missing_ok=True)
    async def guarded_events():
        with operation_lease(session_id, "tool_session", "report_generation"):
            async for event in events():
                yield event

    return StreamingResponse(guarded_events(), media_type="text/event-stream", headers={"Cache-Control": "no-store", "X-Accel-Buffering": "no"})


@router.post("/sessions/{session_id}/download")
async def download_live_docx(session_id: str, request: Request):
    directory = _owned(session_id, request) / "accuracy_report"
    result_file = directory / "result.json"
    auth_header = request.headers.get("authorization", "")
    if not auth_header.startswith("Bearer "):
        raise HTTPException(401, "Google permission is required to download the document")
    if not result_file.is_file():
        raise HTTPException(404, "Generated document not found")
    result = json.loads(result_file.read_text(encoding="utf-8"))
    try:
        async with httpx.AsyncClient(timeout=60, follow_redirects=False) as client:
            headers = {"Authorization": auth_header, "Accept": "application/json"}
            response = await client.post(GOOGLE_DRIVE_DOWNLOAD.format(file_id=quote(result["doc_id"], safe="")), headers=headers)
            if response.status_code >= 400:
                raise RuntimeError()
            operation = response.json()
            for attempt in range(32):
                if operation.get("done") is True:
                    break
                if operation.get("error"):
                    raise RuntimeError()
                operation_name = operation.get("name")
                if not isinstance(operation_name, str) or not operation_name.startswith("operations/"):
                    raise RuntimeError()
                await asyncio.sleep(min(0.5 * (2 ** min(attempt, 3)), 4))
                operation_url = GOOGLE_DRIVE_OPERATIONS.format(quote(operation_name.split("/", 1)[1], safe=""))
                response = await client.get(operation_url, headers=headers)
                if response.status_code >= 400:
                    raise RuntimeError()
                operation = response.json()
            else:
                raise RuntimeError()
            if operation.get("done") is not True or operation.get("error"):
                raise RuntimeError()
            download_url = operation.get("response", {}).get("downloadUri")
            parsed = urlsplit(download_url or "")
            trusted_hosts = {"drive.google.com", "www.googleapis.com", "drive.usercontent.google.com"}
            if parsed.scheme != "https" or not parsed.hostname or not (parsed.hostname in trusted_hosts or parsed.hostname.endswith(".googleusercontent.com")):
                raise RuntimeError()
            response = await client.get(download_url)
            for _ in range(3):
                if response.status_code not in {301, 302, 303, 307, 308}:
                    break
                redirect_url = response.headers.get("location", "")
                parsed = urlsplit(redirect_url)
                if parsed.scheme != "https" or not parsed.hostname or not (parsed.hostname in trusted_hosts or parsed.hostname.endswith(".googleusercontent.com")):
                    raise RuntimeError()
                response = await client.get(redirect_url)
            if response.status_code >= 400 or response.status_code in {301, 302, 303, 307, 308}:
                raise RuntimeError()
    except Exception:
        raise HTTPException(502, "Could not export the current Google Doc")
    from fastapi.responses import Response
    return Response(response.content, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": 'attachment; filename="accuracy-report.docx"', "Cache-Control": "no-store"})
