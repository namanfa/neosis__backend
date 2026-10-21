"""Local-first APIs for Testing Images and Packshots."""
import asyncio
import csv
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
import hashlib
import io
import json
import logging
import os
import re
import tempfile
import threading
import time
import unicodedata
import zipfile
from pathlib import PurePosixPath
from urllib.parse import quote
import xml.etree.ElementTree as ET

from fastapi import APIRouter, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from PIL import Image, ImageCms, ImageOps, UnidentifiedImageError

from storage.brands import add_brand, get_brands

router = APIRouter(prefix="/api/tools", tags=["Internal Tools"])
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
MAX_FILES_PER_BATCH = 16
MAX_PARALLEL_OBJECT_WRITES = 6
MAX_PARALLEL_LIST_CALLS = 8
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_BATCH_BYTES = 100 * 1024 * 1024
MAX_BATCH_UPLOAD_BYTES = 60 * 1024 * 1024
MAX_SHEET_BYTES = 20 * 1024 * 1024
MAX_LOGO_BYTES = 1024 * 1024
IMAGE_DERIVATIVE_MIN_BYTES = max(0, int(os.environ.get("IMAGE_DERIVATIVE_MIN_BYTES", 3 * 1024 * 1024)))
IMAGE_HEADER_BYTES = 524288
MANIFEST = "manifest.json"
BRAND_INDEX_KEY = "_index/brands.json"
GALLERY_INDEX_PREFIX = "_index/galleries"
THUMB_PREFIX = "thumbs"
LOGO_PREFIX = "logos"
logger = logging.getLogger(__name__)
_index_lock = threading.RLock()
_derivative_locks = {}
_derivative_locks_guard = threading.Lock()
_derivative_pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix="image-derivative")
_brand_data_cache = {}
_brand_data_lock = threading.Lock()
_brand_index_cache = {}
_gallery_data_cache = {}
_packshot_label_locks = {}
_packshot_label_locks_guard = threading.Lock()


def _storage_cache_key(storage):
    root = getattr(storage, "root", None)
    if root is not None:
        return (type(storage).__qualname__, str(root))
    return (type(storage).__qualname__, getattr(storage, "endpoint_url", ""), getattr(storage, "bucket", ""))


def _store_objects(storage, objects):
    """Write independent objects concurrently; object storage writes are network-bound."""
    failures = {}
    with ThreadPoolExecutor(max_workers=MAX_PARALLEL_OBJECT_WRITES) as pool:
        futures = {
            pool.submit(storage.put, key, data, content_type, "private, max-age=86400"): (key, filename)
            for key, data, content_type, filename in objects
        }
        for future in as_completed(futures):
            key, filename = futures[future]
            try:
                future.result()
            except Exception as exc:
                failures[key] = {"file": filename, "reason": str(exc)}
    return failures


def _response_timing(response, started, list_ms=0.0, storage_calls=0, signing_ms=0.0):
    total_ms = (time.perf_counter() - started) * 1000
    response.headers["Server-Timing"] = (
        f"total;dur={total_ms:.1f}, storage-list;dur={list_ms:.1f}, "
        f"storage-calls;desc=\"{storage_calls}\", url-signing;dur={signing_ms:.1f}"
    )
    return response


def _etag_json(data):
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return '"' + hashlib.sha256(payload).hexdigest() + '"'


def _json_response(request, data, started, list_ms=0.0, storage_calls=0, signing_ms=0.0):
    etag = _etag_json(data)
    headers = {"ETag": etag, "Cache-Control": "private, max-age=0, must-revalidate"}
    timing_response = Response(status_code=304, headers=headers) if request.headers.get("if-none-match") == etag else JSONResponse(data, headers=headers)
    return _response_timing(timing_response, started, list_ms, storage_calls, signing_ms)


def _brand_index(storage):
    """Read the durable count index, rebuilding it once when missing or outdated."""
    storage_id = _storage_cache_key(storage)
    now = time.monotonic()
    with _brand_data_lock:
        cached = _brand_index_cache.get(storage_id)
        if cached and cached[0] > now:
            return cached[1], cached[2]
    brands = _cached_brands(storage)
    try:
        index = storage.get_json(BRAND_INDEX_KEY)
        indexed = index.get("brands", {}) if isinstance(index, dict) and index.get("version") == 2 else {}
        if set(indexed) == {item.get("slug", "") for item in brands}:
            with _brand_data_lock:
                _brand_index_cache[storage_id] = (now + 15, brands, indexed)
            return brands, indexed
    except Exception:
        pass
    with _index_lock:
        try:
            index = storage.get_json(BRAND_INDEX_KEY)
            indexed = index.get("brands", {}) if isinstance(index, dict) and index.get("version") == 2 else {}
            if set(indexed) == {item.get("slug", "") for item in brands}:
                with _brand_data_lock:
                    _brand_index_cache[storage_id] = (now + 15, brands, indexed)
                return brands, indexed
        except Exception:
            indexed = {}
        refreshed = _scan_brand_entries(storage, brands)
        storage.put_json(BRAND_INDEX_KEY, {"version": 2, "updated_at": datetime.now(timezone.utc).isoformat(), "brands": refreshed})
        with _brand_data_lock:
            _brand_index_cache[storage_id] = (time.monotonic() + 15, brands, refreshed)
        return brands, refreshed


def _cached_brands(storage):
    storage_id = _storage_cache_key(storage)
    now = time.monotonic()
    with _brand_data_lock:
        cached = _brand_data_cache.get(storage_id)
        if cached and cached[0] > now:
            return cached[1]
    brands = get_brands(storage)
    with _brand_data_lock:
        _brand_data_cache[storage_id] = (now + 15, brands)
    return brands


def _scan_brand_entries(storage, brands):
    pairs = {}
    with ThreadPoolExecutor(max_workers=MAX_PARALLEL_LIST_CALLS) as pool:
        jobs = {}
        for brand in brands:
            slug = brand.get("slug", "")
            jobs[pool.submit(storage.list_entries, f"testing/{slug}")] = (slug, "testing")
            jobs[pool.submit(storage.list_entries, f"packshots/{slug}")] = (slug, "packshots")
        for future in as_completed(jobs):
            slug, collection = jobs[future]
            pairs.setdefault(slug, {})[collection] = future.result()
    result = {}
    for brand in brands:
        slug = brand.get("slug", "")
        testing = pairs.get(slug, {}).get("testing", [])
        packshots = [entry for entry in pairs.get(slug, {}).get("packshots", []) if entry["key"].rsplit("/", 1)[-1] != MANIFEST]
        result[slug] = {
            "testing_count": len(testing),
            "testing_assets": sorted({entry["key"].split("/")[2] for entry in testing if len(entry["key"].split("/")) > 3}),
            "packshot_count": len(packshots),
            "logo": brand.get("logo"),
        }
    return result


def _refresh_brand_index(storage, brand_slug):
    brands = get_brands(storage)
    brand = next((item for item in brands if item.get("slug") == brand_slug), None)
    if brand is None:
        return
    scanned = _scan_brand_entries(storage, [brand])
    with _index_lock:
        try:
            current = storage.get_json(BRAND_INDEX_KEY).get("brands", {})
        except Exception:
            current = {}
        current[brand_slug] = scanned[brand_slug]
        storage.put_json(BRAND_INDEX_KEY, {"version": 2, "updated_at": datetime.now(timezone.utc).isoformat(), "brands": current})
        with _brand_data_lock:
            _brand_index_cache[_storage_cache_key(storage)] = (time.monotonic() + 15, brands, current)
            _brand_data_cache[_storage_cache_key(storage)] = (time.monotonic() + 15, brands)


def _save_brand_logo_metadata(storage, brand_slug, logo_record):
    with _index_lock:
        brands = get_brands(storage)
        for brand in brands:
            if brand.get("slug") == brand_slug:
                brand["logo"] = logo_record
                break
        storage.put_json("brands.json", brands)
        with _brand_data_lock:
            _brand_data_cache[_storage_cache_key(storage)] = (time.monotonic() + 15, brands)


def _gallery_index_key(kind, slug):
    return f"{GALLERY_INDEX_PREFIX}/{kind}/{slug}.json"


def _scan_gallery(storage, kind, slug, entries=None, dimensions=None, derivative_decisions=None, previous=None):
    prefix = f"{kind}/{slug}"
    entries = entries if entries is not None else storage.list_entries(prefix)
    labels = {}
    if kind == "packshots":
        try:
            labels = {row.get("file"): row for row in _manifest(storage, f"{prefix}/{MANIFEST}") if isinstance(row, dict)}
        except Exception:
            labels = {}
    records = []
    assets = set()
    for entry in entries:
        key = entry["key"]
        name = key.rsplit("/", 1)[-1]
        if kind == "packshots" and name == MANIFEST:
            continue
        if os.path.splitext(name)[1].lower() not in IMAGE_EXTENSIONS:
            continue
        parts = key.split("/")
        asset = "general" if kind == "testing" and len(parts) == 3 else (parts[2] if kind == "testing" and len(parts) > 3 else "")
        if asset and asset != "general":
            assets.add(asset)
        version = str(entry.get("etag") or f"{entry.get('last_modified') or 0}-{entry.get('size') or 0}").strip('"')
        label = labels.get(name, {})
        known_dimensions = (dimensions or {}).get(key, {})
        old_record = (previous or {}).get(key, {})
        derivative = (derivative_decisions or {}).get(key, label.get("derivative", old_record.get("derivative")))
        record = {
            "key": key, "file": name, "asset": asset, "size": int(entry.get("size") or 0), "etag": version,
            "width": known_dimensions.get("width"), "height": known_dimensions.get("height"),
            "annotation_name": label.get("annotation_name", os.path.splitext(name)[0]),
            "name": label.get("name", ""), "id": label.get("id", ""),
            "updated_by": label.get("updated_by"), "updated_at": label.get("updated_at"),
        }
        record["width"] = record["width"] or old_record.get("width")
        record["height"] = record["height"] or old_record.get("height")
        record["derivative"] = derivative
        record["label_etag"] = _label_etag(record, name)
        records.append(record)
    return {"version": 2, "images": records, "assets": sorted(assets)}


def _gallery_manifest(storage, kind, slug, bypass_memory=False):
    memory_key = (_storage_cache_key(storage), kind, slug)
    cached = _gallery_data_cache.get(memory_key)
    if not bypass_memory and cached and cached[0] > time.monotonic():
        return cached[1], False
    cache_key = _gallery_index_key(kind, slug)
    try:
        data = storage.get_json(cache_key)
        if isinstance(data, dict) and data.get("version") == 2 and isinstance(data.get("images"), list) and isinstance(data.get("assets"), list):
            _gallery_data_cache[memory_key] = (time.monotonic() + 15, data)
            return data, False
    except Exception:
        pass
    data = _scan_gallery(storage, kind, slug)
    storage.put_json(cache_key, data)
    _gallery_data_cache[memory_key] = (time.monotonic() + 15, data)
    return data, True


def _refresh_gallery(storage, kind, slug, entries=None, dimensions=None, derivative_decisions=None):
    with _index_lock:
        previous = {}
        memory = _gallery_data_cache.get((_storage_cache_key(storage), kind, slug))
        if memory:
            previous = {row["key"]: row for row in memory[1].get("images", [])}
        else:
            try:
                previous = {row["key"]: row for row in storage.get_json(_gallery_index_key(kind, slug)).get("images", [])}
            except Exception:
                previous = {}
        data = _scan_gallery(storage, kind, slug, entries, dimensions, derivative_decisions, previous)
        storage.put_json(_gallery_index_key(kind, slug), data)
        _gallery_data_cache[(_storage_cache_key(storage), kind, slug)] = (time.monotonic() + 15, data)
        _refresh_brand_index(storage, slug)


def _version_for(entry):
    value = str(entry.get("etag") or f"{entry.get('last_modified') or 0}-{entry.get('size') or 0}").strip('"')
    return re.sub(r"[^A-Za-z0-9._-]", "_", value) or "0"


def _image_url(storage, kind, slug, entry, variant):
    key = entry["key"]
    if entry.get("derivative") == "none":
        return storage.url_for_display(key)
    version = quote(_version_for(entry), safe="")
    relative = key.split("/", 2)[2]
    return f"/api/tools/thumb/{kind}/{quote(slug, safe='')}/{quote(relative, safe='/')}?v={version}&size={variant}"


def _image_content_type(key):
    ext = os.path.splitext(key)[1].lower()
    return {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}.get(ext, "application/octet-stream")


def _derivative_key(kind, slug, source_key, version, variant):
    source_id = hashlib.sha256(source_key.encode("utf-8")).hexdigest()[:20]
    version_id = re.sub(r"[^A-Za-z0-9._-]", "_", version)[:80]
    return f"{THUMB_PREFIX}/{kind}/{slug}/{source_id}-{version_id}-{variant}.webp"


def _image_dimensions(data):
    with Image.open(io.BytesIO(data)) as source:
        width, height = source.size
        if source.getexif().get(274) in {5, 6, 7, 8}:
            width, height = height, width
        return int(width), int(height)


def _derivative_gate(size, width, height):
    try:
        threshold = max(0, int(os.environ.get("IMAGE_DERIVATIVE_MIN_BYTES", IMAGE_DERIVATIVE_MIN_BYTES)))
    except (TypeError, ValueError):
        threshold = IMAGE_DERIVATIVE_MIN_BYTES
    return int(size) >= threshold or max(int(width), int(height)) > 1920


def _existing_derivative_result(storage, kind, slug, source_key, version, original_size, width, height):
    card = _derivative_key(kind, slug, source_key, version, "card")
    preview = _derivative_key(kind, slug, source_key, version, "preview")
    card_size = storage.get_size(card)
    preview_size = storage.get_size(preview)
    if card_size is None and preview_size is None:
        return None
    use_card = card_size is not None and card_size < original_size
    use_preview = preview_size is not None and preview_size < original_size
    result = {
        "card": card if use_card else source_key,
        "preview": preview if use_preview else source_key,
        "width": width,
        "height": height,
        "original": not (use_card or use_preview),
        "derivative": "generated" if use_card or use_preview else "none",
        "card_bytes": card_size if use_card else 0,
        "preview_bytes": preview_size if use_preview else 0,
        "original_bytes": original_size,
        "storage_calls": 2,
    }
    return result


def _label_etag(record, filename):
    payload = {
        "file": filename,
        "annotation_name": record.get("annotation_name") or os.path.splitext(filename)[0],
        "name": record.get("name") or "",
        "id": record.get("id") or "",
        "updated_at": record.get("updated_at") or "",
        "updated_by": record.get("updated_by") or "",
    }
    digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()[:24]
    return f'"{digest}"'


def _normalize_label_value(field, value, filename):
    if field not in {"annotation_name", "name", "id"}:
        raise HTTPException(400, "Field must be annotation, name, or id")
    if not isinstance(value, str):
        raise HTTPException(400, "Label value must be text")
    if any(unicodedata.category(character) == "Cc" for character in value):
        raise HTTPException(400, "Label cannot contain control characters")
    normalized = re.sub(r"\s+", " ", value.strip())
    if len(normalized) > 200:
        raise HTTPException(400, "Label must be 200 characters or fewer")
    if field == "annotation_name" and not normalized:
        normalized = os.path.splitext(filename)[0]
    return normalized


def _make_webp(data, max_side, quality):
    Image.MAX_IMAGE_PIXELS = min(Image.MAX_IMAGE_PIXELS or 89_478_485, 178_956_970)
    with Image.open(io.BytesIO(data)) as source:
        source.load()
        oriented = ImageOps.exif_transpose(source)
        width, height = oriented.size
        if max(width, height) <= max_side and len(data) < 200 * 1024:
            return None, width, height
        oriented.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
        has_alpha = "A" in oriented.getbands() or "transparency" in oriented.info
        alpha = oriented.getchannel("A") if "A" in oriented.getbands() else None
        converted = oriented.convert("RGB")
        profile_data = source.info.get("icc_profile")
        if profile_data:
            try:
                source_profile = ImageCms.ImageCmsProfile(io.BytesIO(profile_data))
                converted = ImageCms.profileToProfile(converted, source_profile, ImageCms.createProfile("sRGB"), outputMode="RGB")
            except Exception:
                pass
        if has_alpha:
            converted.putalpha(alpha if alpha is not None else oriented.convert("RGBA").getchannel("A"))
        output = io.BytesIO()
        converted.save(output, format="WEBP", quality=quality, method=4)
        return output.getvalue(), width, height


def _generate_derivatives(storage, kind, slug, source_key, version, data=None):
    card_key = _derivative_key(kind, slug, source_key, version, "card")
    preview_key = _derivative_key(kind, slug, source_key, version, "preview")
    had_data = data is not None
    known = _gallery_record_for(storage, kind, slug, source_key)
    if data is None and known and known.get("derivative") == "none":
        return {"card": source_key, "preview": source_key, "width": known.get("width"), "height": known.get("height"), "original": True, "derivative": "none", "card_bytes": 0, "preview_bytes": 0, "original_bytes": known.get("size", 0), "storage_calls": 0}
    if data is None and known and known.get("size") is not None and known.get("width") and known.get("height"):
        if not _derivative_gate(known["size"], known["width"], known["height"]):
            result = {"card": source_key, "preview": source_key, "width": known["width"], "height": known["height"], "original": True, "derivative": "none", "card_bytes": 0, "preview_bytes": 0, "original_bytes": known["size"], "storage_calls": 0}
            _record_derivative_decision(storage, kind, slug, source_key, version, result)
            return result
        existing = _existing_derivative_result(storage, kind, slug, source_key, version, known["size"], known["width"], known["height"])
        if existing:
            _record_image_dimensions(storage, kind, slug, source_key, version, known["width"], known["height"])
            _record_derivative_decision(storage, kind, slug, source_key, version, existing)
            return existing
    if data is None:
        with storage.get(source_key) as source:
            data = source.read(MAX_IMAGE_BYTES + 1)
        storage_calls = 1
    else:
        storage_calls = 0
    if len(data) > MAX_IMAGE_BYTES:
        raise ValueError("Stored image exceeds the 20 MB image limit")
    width, height = _image_dimensions(data)
    if not _derivative_gate(len(data), width, height):
        result = {"card": source_key, "preview": source_key, "width": width, "height": height, "original": True, "derivative": "none", "card_bytes": 0, "preview_bytes": 0, "original_bytes": len(data), "original_data": data, "storage_calls": storage_calls}
        _record_derivative_decision(storage, kind, slug, source_key, version, result)
        return result
    if not had_data:
        existing = _existing_derivative_result(storage, kind, slug, source_key, version, len(data), width, height)
        if existing:
            _record_derivative_decision(storage, kind, slug, source_key, version, existing)
            return existing
    card, width, height = _make_webp(data, 640, 80)
    preview, _, _ = _make_webp(data, 1920, 85)
    cache = "public, max-age=31536000, immutable"
    if card is not None and len(card) < len(data):
        storage.put(card_key, card, "image/webp", cache_control=cache)
        storage_calls += 1
    else:
        card_key = source_key
    if preview is not None and len(preview) < len(data):
        storage.put(preview_key, preview, "image/webp", cache_control=cache)
        storage_calls += 1
        preview_bytes = len(preview)
    else:
        preview_key = source_key
        preview_bytes = 0
    card_bytes = len(card) if card is not None and card_key != source_key else 0
    generated = card_key != source_key or preview_key != source_key
    result = {"card": card_key, "preview": preview_key, "width": width, "height": height, "original": not generated, "derivative": "generated" if generated else "none", "card_bytes": card_bytes, "preview_bytes": preview_bytes, "original_bytes": len(data), "storage_calls": storage_calls}
    _record_derivative_decision(storage, kind, slug, source_key, version, result)
    return result


def _record_image_dimensions(storage, kind, slug, source_key, version, width, height):
    if not width or not height:
        return
    cache_key = (_storage_cache_key(storage), kind, slug)
    with _index_lock:
        cached = _gallery_data_cache.get(cache_key)
        try:
            data = cached[1] if cached else storage.get_json(_gallery_index_key(kind, slug))
        except Exception:
            return
        changed = False
        for entry in data.get("images", []):
            if entry.get("key") == source_key and _version_for(entry) == version and not entry.get("width"):
                entry["width"], entry["height"] = int(width), int(height)
                changed = True
                break
        if changed:
            storage.put_json(_gallery_index_key(kind, slug), data)
            _gallery_data_cache[cache_key] = (time.monotonic() + 15, data)


def _gallery_record_for(storage, kind, slug, source_key):
    cache_key = (_storage_cache_key(storage), kind, slug)
    cached = _gallery_data_cache.get(cache_key)
    if cached and cached[0] > time.monotonic():
        return next((row for row in cached[1].get("images", []) if row.get("key") == source_key), None)
    try:
        data = storage.get_json(_gallery_index_key(kind, slug))
    except Exception:
        return None
    _gallery_data_cache[cache_key] = (time.monotonic() + 15, data)
    return next((row for row in data.get("images", []) if row.get("key") == source_key), None)


def _record_derivative_decision(storage, kind, slug, source_key, version, result):
    cache_key = (_storage_cache_key(storage), kind, slug)
    with _index_lock:
        cached = _gallery_data_cache.get(cache_key)
        try:
            data = cached[1] if cached else storage.get_json(_gallery_index_key(kind, slug))
        except Exception:
            try:
                data = _scan_gallery(storage, kind, slug)
            except Exception:
                data = {"version": 2, "images": [], "assets": []}
        record = next((row for row in data.get("images", []) if row.get("key") == source_key and row.get("etag") == version), None)
        if record is None:
            record = next((row for row in data.get("images", []) if row.get("key") == source_key), None)
        if record is None:
            return
        record.update({"derivative": result["derivative"], "width": result["width"], "height": result["height"]})
        record["label_etag"] = _label_etag(record, record["file"])
        storage.put_json(_gallery_index_key(kind, slug), data)
        _gallery_data_cache[cache_key] = (time.monotonic() + 15, data)


def _resolve_gallery_derivatives(storage, kind, slug, gallery):
    changed = False
    if kind == "packshots":
        for row in gallery.get("images", []):
            if not row.get("label_etag"):
                row["label_etag"] = _label_etag(row, row.get("file", ""))
                changed = True
    unresolved = [row for row in gallery.get("images", []) if row.get("derivative") not in {"none", "generated"}]
    if not unresolved:
        if changed:
            storage.put_json(_gallery_index_key(kind, slug), gallery)
            _gallery_data_cache[(_storage_cache_key(storage), kind, slug)] = (time.monotonic() + 15, gallery)
        return gallery

    def inspect(row):
        try:
            header = storage.get_prefix(row["key"], IMAGE_HEADER_BYTES)
            width, height = _image_dimensions(header)
        except Exception:
            with storage.get(row["key"]) as source:
                data = source.read(MAX_IMAGE_BYTES + 1)
            width, height = _image_dimensions(data)
        return row["key"], width, height, "generated" if _derivative_gate(row.get("size", 0), width, height) else "none"

    resolved = {}
    with ThreadPoolExecutor(max_workers=min(MAX_PARALLEL_LIST_CALLS, len(unresolved))) as pool:
        futures = {pool.submit(inspect, row): row for row in unresolved}
        for future in as_completed(futures):
            try:
                key, width, height, decision = future.result()
                resolved[key] = (width, height, decision)
            except Exception:
                # A corrupt or unsupported image should remain viewable through its original URL.
                row = futures[future]
                resolved[row["key"]] = (row.get("width"), row.get("height"), "none")
    for row in gallery.get("images", []):
        if row.get("key") in resolved:
            row["width"], row["height"], row["derivative"] = resolved[row["key"]]
            row["label_etag"] = _label_etag(row, row["file"])
    with _index_lock:
        storage.put_json(_gallery_index_key(kind, slug), gallery)
        _gallery_data_cache[(_storage_cache_key(storage), kind, slug)] = (time.monotonic() + 15, gallery)
    return gallery


def _ensure_derivatives(storage, kind, slug, source_key, version, data=None):
    lock_key = (kind, slug, source_key, version)
    with _derivative_locks_guard:
        lock = _derivative_locks.setdefault(lock_key, threading.Lock())
    with lock:
        return _generate_derivatives(storage, kind, slug, source_key, version, data)


def _schedule_derivatives(storage, kind, slug, source_key, version, data):
    def work():
        try:
            _ensure_derivatives(storage, kind, slug, source_key, version, data)
        except Exception as exc:
            logger.warning("Could not generate %s image derivatives for %s: %s", kind, source_key, exc)
    _derivative_pool.submit(work)


def _prepare_logo(upload_name, data):
    ext = os.path.splitext(upload_name or "")[1].lower()
    if ext not in {".png", ".jpg", ".jpeg", ".webp", ".svg"}:
        raise HTTPException(400, "Logo must be PNG, JPG, WebP, or SVG")
    if not data or len(data) > MAX_LOGO_BYTES:
        raise HTTPException(413, "Logo must be smaller than 1 MB")
    if ext == ".svg":
        if b"<!doctype" in data.lower() or b"<!entity" in data.lower():
            raise HTTPException(400, "SVG logo contains unsupported declarations")
        try:
            root = ET.fromstring(data)
        except ET.ParseError:
            raise HTTPException(400, "Invalid SVG logo")
        for element in root.iter():
            tag = element.tag.rsplit("}", 1)[-1].casefold()
            if tag in {"script", "foreignobject", "iframe", "object", "embed"}:
                raise HTTPException(400, "SVG logo contains unsupported content")
            for name, value in element.attrib.items():
                attr = name.rsplit("}", 1)[-1].casefold()
                if attr.startswith("on") or (attr in {"href", "src"} and value and not value.startswith("#")) or "url(" in value.casefold():
                    raise HTTPException(400, "SVG logo cannot reference scripts or external resources")
        return data, ext, "image/svg+xml"
    try:
        Image.MAX_IMAGE_PIXELS = min(Image.MAX_IMAGE_PIXELS or 89_478_485, 178_956_970)
        with Image.open(io.BytesIO(data)) as image:
            if image.format not in {"PNG", "JPEG", "WEBP"}:
                raise HTTPException(400, "Logo content does not match a supported image")
            image.load()
            image = ImageOps.exif_transpose(image)
            image.thumbnail((128, 128), Image.Resampling.LANCZOS)
            image = image.convert("RGBA" if "A" in image.getbands() or "transparency" in image.info else "RGB")
            output = io.BytesIO()
            image.save(output, format="PNG", optimize=True)
            return output.getvalue(), ".png", "image/png"
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError):
        raise HTTPException(400, "Invalid or unsafe logo image")


def _save_logo(storage, brand, data, ext, content_type):
    digest = hashlib.sha256(data).hexdigest()[:24]
    key = f"{LOGO_PREFIX}/{brand['slug']}/{digest}{ext}"
    storage.put(key, data, content_type, cache_control="public, max-age=31536000, immutable")
    return {"key": key, "version": digest, "content_type": content_type}


def _stream_storage_object(stream):
    try:
        while chunk := stream.read(64 * 1024):
            yield chunk
    finally:
        stream.close()


def _brand(storage, slug):
    brand = next((item for item in _cached_brands(storage) if item.get("slug") == slug), None)
    if brand is None:
        raise HTTPException(404, "Brand not found")
    return brand


def _slug(value, label="Name"):
    if not isinstance(value, str) or "\x00" in value or "\\" in value:
        raise HTTPException(400, f"Invalid {label.lower()}")
    value = " ".join(value.split())
    if not value or value in {".", ".."} or value.startswith("/"):
        raise HTTPException(400, f"{label} cannot be empty or contain a path")
    result = re.sub(r"[^\w-]+", "-", value.casefold(), flags=re.UNICODE).strip("-_")
    if not result or result in {".", ".."}:
        raise HTTPException(400, f"Invalid {label.lower()}")
    return result


def _filename(value):
    if not isinstance(value, str) or "\x00" in value or "\\" in value:
        raise HTTPException(400, "Invalid image path")
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in value.split("/")):
        raise HTTPException(400, "Invalid image path")
    name = path.name
    ext = os.path.splitext(name)[1].lower()
    if ext not in IMAGE_EXTENSIONS:
        return None
    stem = os.path.splitext(name)[0]
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-") or "image"
    return stem + ext


def _is_image(data, ext):
    if ext in {".jpg", ".jpeg"}:
        return data.startswith(b"\xff\xd8\xff")
    if ext == ".png":
        return data.startswith(b"\x89PNG\r\n\x1a\n")
    return len(data) >= 12 and data.startswith(b"RIFF") and data[8:12] == b"WEBP"


async def _read_image(upload):
    ext = os.path.splitext(upload.filename or "")[1].lower()
    if ext not in IMAGE_EXTENSIONS:
        return None
    content = await upload.read(MAX_IMAGE_BYTES + 1)
    if len(content) > MAX_IMAGE_BYTES:
        raise HTTPException(413, f"{upload.filename}: image exceeds the 20 MB limit")
    if not content or not _is_image(content, ext):
        raise HTTPException(400, f"{upload.filename}: file content does not match a supported image")
    return content


def _manifest(storage, key):
    if not storage.exists(key):
        return []
    data = storage.get_json(key)
    if not isinstance(data, list):
        raise HTTPException(500, "Stored manifest is invalid")
    return data


def _unique_key(storage, prefix, filename, reserved=None, existing=None):
    reserved = reserved if reserved is not None else set()
    stem, ext = os.path.splitext(filename)
    candidate = filename
    suffix = 1
    key = f"{prefix}/{candidate}"
    while key in reserved or (key in existing if existing is not None else storage.exists(key)):
        candidate = f"{stem}_{suffix}{ext}"
        suffix += 1
        key = f"{prefix}/{candidate}"
    reserved.add(key)
    return candidate


@router.get("/brands")
def list_tool_brands(request: Request):
    started = time.perf_counter()
    storage = request.app.state.storage
    cache = _brand_index_cache.get(_storage_cache_key(storage))
    cache_hit = bool(cache and cache[0] > time.monotonic())
    list_started = time.perf_counter()
    brands, indexed = _brand_index(storage)
    list_ms = (time.perf_counter() - list_started) * 1000 if not cache_hit else 0
    storage_calls = 0 if cache_hit else 4 + 2 * len(brands)
    result = []
    signing_started = time.perf_counter()
    for brand in brands:
        slug = brand.get("slug", "")
        inventory = indexed.get(slug, {})
        logo = inventory.get("logo") or brand.get("logo") or {}
        result.append({
            **brand,
            "testing_count": inventory.get("testing_count", 0),
            "testing_assets": inventory.get("testing_assets", []),
            "packshot_count": inventory.get("packshot_count", 0),
            "logo_url": f"/api/tools/logos/{quote(slug, safe='')}/{quote(logo.get('version', ''), safe='')}{os.path.splitext(logo.get('key', ''))[1]}" if logo.get("key") else None,
        })
    signing_ms = (time.perf_counter() - signing_started) * 1000
    return _json_response(request, {"brands": result}, started, list_ms, storage_calls, signing_ms)


@router.post("/brands")
async def create_tool_brand(request: Request, name: str = Form(...), logo: UploadFile | None = File(None)):
    try:
        prepared_logo = None
        if logo and logo.filename:
            prepared_logo = _prepare_logo(logo.filename, await logo.read(MAX_LOGO_BYTES + 1))
        storage = request.app.state.storage
        brand = add_brand(storage, name)
    except ValueError as exc:
        raise HTTPException(409 if "already exists" in str(exc) or "conflicts" in str(exc) else 400, str(exc))
    warning = None
    if prepared_logo:
        try:
            brand["logo"] = _save_logo(storage, brand, *prepared_logo)
            _save_brand_logo_metadata(storage, brand["slug"], brand["logo"])
        except Exception as exc:
            logger.warning("Brand created but logo storage failed for %s: %s", brand["slug"], exc)
            warning = "Brand created, but its logo could not be saved. You can add it again later."
    _refresh_brand_index(storage, brand["slug"])
    with _brand_data_lock:
        _brand_data_cache[_storage_cache_key(storage)] = (time.monotonic() + 15, get_brands(storage))
    return {**brand, "logo_url": f"/api/tools/logos/{brand['slug']}/{brand['logo']['version']}{os.path.splitext(brand['logo']['key'])[1]}" if brand.get("logo") else None, "logo_warning": warning}


@router.post("/brands/{brand_slug}/logo")
async def set_tool_brand_logo(brand_slug: str, request: Request, logo: UploadFile = File(...)):
    storage = request.app.state.storage
    brand = _brand(storage, brand_slug)
    prepared = _prepare_logo(logo.filename or "", await logo.read(MAX_LOGO_BYTES + 1))
    record = _save_logo(storage, brand, *prepared)
    _save_brand_logo_metadata(storage, brand_slug, record)
    _refresh_brand_index(storage, brand_slug)
    with _brand_data_lock:
        cached = _brand_index_cache.get(_storage_cache_key(storage))
        if cached:
            index = dict(cached[2])
            index[brand_slug] = {**index.get(brand_slug, {}), "logo": record}
            _brand_index_cache[_storage_cache_key(storage)] = (time.monotonic() + 15, cached[1], index)
    return {"logo_url": f"/api/tools/logos/{brand_slug}/{record['version']}{os.path.splitext(record['key'])[1]}", "logo_warning": None}


@router.get("/logos/{brand_slug}/{version}")
def get_tool_brand_logo(brand_slug: str, version: str, request: Request):
    _brand(request.app.state.storage, brand_slug)
    match = re.fullmatch(r"([a-f0-9]{24})(\.png|\.svg)", version)
    if not match:
        raise HTTPException(404, "Brand logo not found")
    digest, ext = match.groups()
    key = f"{LOGO_PREFIX}/{brand_slug}/{digest}{ext}"
    storage = request.app.state.storage
    try:
        stream = storage.get(key)
    except Exception:
        raise HTTPException(404, "Brand logo not found")
    content_type = "image/svg+xml" if ext == ".svg" else "image/png"
    headers = {"Cache-Control": "public, max-age=31536000, immutable", "ETag": f'"{digest}"', "X-Content-Type-Options": "nosniff"}
    if ext == ".svg":
        headers["Content-Security-Policy"] = "sandbox"
    return StreamingResponse(_stream_storage_object(stream), media_type=content_type, headers=headers)


@router.get("/testing/{brand_slug}/images")
def list_testing_images(brand_slug: str, request: Request, asset: str = "all"):
    started = time.perf_counter()
    storage = request.app.state.storage
    _brand(storage, brand_slug)
    list_started = time.perf_counter()
    manifest, rebuilt = _gallery_manifest(storage, "testing", brand_slug)
    manifest = _resolve_gallery_derivatives(storage, "testing", brand_slug, manifest)
    list_ms = (time.perf_counter() - list_started) * 1000 if rebuilt else 0
    images = []
    signing_started = time.perf_counter()
    for entry in manifest["images"]:
        if asset != "all" and entry["asset"] != asset:
            continue
        image = {
            **entry,
            "url": storage.url_for_display(entry["key"]),
            "thumb_url": _image_url(storage, "testing", brand_slug, entry, "card"),
            "preview_url": _image_url(storage, "testing", brand_slug, entry, "preview"),
        }
        images.append(image)
    signing_ms = (time.perf_counter() - signing_started) * 1000
    return _json_response(request, {"images": images, "assets": manifest["assets"]}, started, list_ms, 1 if rebuilt else 0, signing_ms)


@router.get("/thumb/{kind}/{brand_slug}/{relative_key:path}")
async def get_tool_image_derivative(kind: str, brand_slug: str, relative_key: str, request: Request, v: str, size: str = "card"):
    started = time.perf_counter()
    if kind not in {"testing", "packshots"} or size not in {"card", "preview"}:
        raise HTTPException(404, "Image derivative not found")
    path = PurePosixPath(relative_key)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in relative_key.split("/")) or path.suffix.lower() not in IMAGE_EXTENSIONS:
        raise HTTPException(404, "Image derivative not found")
    _brand(request.app.state.storage, brand_slug)
    source_key = f"{kind}/{brand_slug}/{relative_key}"
    storage = request.app.state.storage
    try:
        result = await asyncio.to_thread(_ensure_derivatives, storage, kind, brand_slug, source_key, v)
        derivative_key = result[size]
        if derivative_key == source_key and result.get("original_data") is not None:
            data = result["original_data"]
        else:
            with storage.get(derivative_key) as source:
                data = source.read()
    except (FileNotFoundError, OSError, ValueError, UnidentifiedImageError):
        raise HTTPException(404, "Image derivative not found")
    etag = f'"{v}-{size}"'
    headers = {"Cache-Control": "public, max-age=31536000, immutable", "ETag": etag, "X-Content-Type-Options": "nosniff"}
    if request.headers.get("if-none-match") == etag:
        response = Response(status_code=304, headers=headers)
    else:
        response = Response(content=data, media_type="image/webp" if derivative_key != source_key else _image_content_type(source_key), headers=headers)
    response = _response_timing(response, started, storage_calls=result.get("storage_calls", 0) + (0 if request.headers.get("if-none-match") == etag else 1))
    return response


@router.post("/testing/{brand_slug}/upload")
async def upload_testing_images(
    brand_slug: str,
    request: Request,
    files: list[UploadFile] = File(...),
    relative_paths: str = Form(...),
):
    storage = request.app.state.storage
    _brand(storage, brand_slug)
    if not 1 <= len(files) <= MAX_FILES_PER_BATCH:
        raise HTTPException(400, f"Upload between 1 and {MAX_FILES_PER_BATCH} files per batch")
    try:
        paths = json.loads(relative_paths)
    except (TypeError, json.JSONDecodeError):
        raise HTTPException(400, "Invalid file path metadata")
    if not isinstance(paths, list) or len(paths) != len(files) or not all(isinstance(path, str) for path in paths):
        raise HTTPException(400, "File path metadata does not match uploaded files")

    added = ignored = failed = total_bytes = 0
    failures = []
    reserved = set()
    existing = set(storage.list(f"testing/{brand_slug}"))
    pending = []
    for upload, relative_path in zip(files, paths):
        try:
            safe_name = _filename(relative_path)
            if safe_name is None:
                ignored += 1
                continue
            if os.path.splitext(upload.filename or "")[1].lower() != os.path.splitext(safe_name)[1].lower():
                raise HTTPException(400, "Image path extension does not match uploaded file")
            parts = relative_path.split("/")
            if any(part in {"", ".", ".."} for part in parts):
                raise HTTPException(400, "Invalid image path")
            asset_slug = _slug(parts[0], "Asset") if len(parts) > 1 else None
            data = await _read_image(upload)
            if data is None:
                ignored += 1
                continue
            total_bytes += len(data)
            if total_bytes > MAX_BATCH_BYTES:
                raise HTTPException(413, "Upload batch exceeds the 100 MB limit")
            prefix = f"testing/{brand_slug}" + (f"/{asset_slug}" if asset_slug else "")
            name = _unique_key(storage, prefix, safe_name, reserved, existing)
            content_type = f"image/{'jpeg' if name.endswith(('.jpg', '.jpeg')) else name.rsplit('.', 1)[-1]}"
            pending.append((f"{prefix}/{name}", data, content_type, upload.filename or name))
        except HTTPException as exc:
            failed += 1
            failures.append({"file": upload.filename or "unknown", "reason": exc.detail})
        except Exception as exc:
            failed += 1
            failures.append({"file": upload.filename or "unknown", "reason": str(exc)})
    write_failures = _store_objects(storage, pending)
    for failure in write_failures.values():
        failed += 1
        failures.append(failure)
    added += len(pending) - len(write_failures)
    entries = storage.list_entries(f"testing/{brand_slug}")
    by_key = {entry["key"]: entry for entry in entries}
    dimension_map = {}
    derivative_decisions = {}
    for key, data, _, _ in pending:
        if key not in write_failures:
            try:
                width, height = _image_dimensions(data)
                dimension_map[key] = {"width": width, "height": height}
                derivative_decisions[key] = "generated" if _derivative_gate(len(data), width, height) else "none"
            except Exception:
                pass
    _refresh_gallery(storage, "testing", brand_slug, entries, dimension_map, derivative_decisions)
    for key, data, _, _ in pending:
        if key not in write_failures:
            _schedule_derivatives(storage, "testing", brand_slug, key, _version_for(by_key.get(key, {"size": len(data)})), data)
    return {"added": added, "ignored": ignored, "failed": failed, "failures": failures}


def _header_key(value):
    return "".join(str(value or "").casefold().split())


def _image_stem(value):
    return os.path.splitext(os.path.basename(str(value or "").strip().replace("\\", "/")))[0].casefold()


def _read_mapping(filename, content):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in {".xlsx", ".csv"}:
        raise HTTPException(400, "Upload a .xlsx or .csv mapping sheet")
    if len(content) > MAX_SHEET_BYTES:
        raise HTTPException(413, "Mapping sheet exceeds the 20 MB limit")
    try:
        if ext == ".csv":
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
    except Exception as exc:
        raise HTTPException(400, f"Could not read mapping sheet: {exc}")
    if not rows:
        raise HTTPException(400, "Mapping sheet is empty")
    headers = {_header_key(value): i for i, value in enumerate(rows[0])}
    required = {"imagename": "image_name", "annotationname": "annotation_name", "name": "name", "id": "id"}
    missing = [key for key in required if key not in headers]
    if missing:
        labels = {"imagename": "Image name", "annotationname": "Annotation name", "name": "Name", "id": "ID"}
        raise HTTPException(400, "Mapping sheet is missing columns: " + ", ".join(labels[key] for key in missing))
    mapped = {}
    duplicates = set()
    for row in rows[1:]:
        image_name = row[headers["imagename"]] if headers["imagename"] < len(row) else ""
        stem = _image_stem(image_name)
        if not stem:
            continue
        if stem in mapped:
            duplicates.add(str(image_name))
        annotation = row[headers["annotationname"]] if headers["annotationname"] < len(row) else None
        name = row[headers["name"]] if headers["name"] < len(row) else None
        image_id = row[headers["id"]] if headers["id"] < len(row) else None
        mapped[stem] = {
            "image_name": str(image_name).strip(),
            "annotation_name": str(annotation).strip() if annotation is not None else "",
            "name": str(name).strip() if name is not None else "",
            "id": str(image_id).strip() if image_id is not None else "",
        }
    if duplicates:
        raise HTTPException(400, "Duplicate image names in mapping sheet: " + ", ".join(sorted(duplicates)))
    return mapped


@router.post("/packshots/prepare")
async def prepare_packshot_mapping(mapping: UploadFile = File(...)):
    content = await mapping.read(MAX_SHEET_BYTES + 1)
    rows = _read_mapping(mapping.filename or "", content)
    return {"rows": rows}


@router.get("/packshots/{brand_slug}")
def list_packshots(brand_slug: str, request: Request):
    started = time.perf_counter()
    storage = request.app.state.storage
    _brand(storage, brand_slug)
    list_started = time.perf_counter()
    gallery, rebuilt = _gallery_manifest(storage, "packshots", brand_slug, bypass_memory=True)
    gallery = _resolve_gallery_derivatives(storage, "packshots", brand_slug, gallery)
    list_ms = (time.perf_counter() - list_started) * 1000 if rebuilt else 0
    images = []
    signing_started = time.perf_counter()
    for entry in gallery["images"]:
        key, name = entry["key"], entry["file"]
        images.append({
            **entry,
            "url": storage.url_for_display(key),
            "thumb_url": _image_url(storage, "packshots", brand_slug, entry, "card"),
            "preview_url": _image_url(storage, "packshots", brand_slug, entry, "preview"),
        })
    signing_ms = (time.perf_counter() - signing_started) * 1000
    return _json_response(request, {"images": images}, started, list_ms, 1 if rebuilt else 0, signing_ms)


def _packshot_label_lock(storage, slug):
    key = (_storage_cache_key(storage), slug)
    with _packshot_label_locks_guard:
        return _packshot_label_locks.setdefault(key, threading.Lock())


@router.patch("/packshots/{brand_slug}/{image_key:path}/labels")
async def update_packshot_label(brand_slug: str, image_key: str, request: Request):
    if not image_key or "/" in image_key or "\\" in image_key or image_key in {".", ".."}:
        raise HTTPException(400, "Invalid packshot image key")
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(400, "Request body must be valid JSON")
    if not isinstance(payload, dict):
        raise HTTPException(400, "Request body must be a JSON object")
    field = {"annotation": "annotation_name", "name": "name", "id": "id"}.get(payload.get("field"))
    if field is None:
        raise HTTPException(400, "Field must be annotation, name, or id")
    storage = request.app.state.storage
    slug = _slug(brand_slug, "Brand")
    expected = request.headers.get("if-match")
    if not expected:
        raise HTTPException(400, "If-Match is required; reload the packshot before editing")
    started = time.perf_counter()
    lock = _packshot_label_lock(storage, slug)
    with lock, _index_lock:
        try:
            gallery = storage.get_json(_gallery_index_key("packshots", slug))
        except Exception:
            raise HTTPException(404, "Packshot gallery not found")
        record = next((row for row in gallery.get("images", []) if row.get("file") == image_key), None)
        if record is None:
            raise HTTPException(404, "Packshot not found")
        filename = record["file"]
        current_etag = record.get("label_etag") or _label_etag(record, filename)
        if expected != current_etag:
            raise HTTPException(409, "This label was changed by someone else")
        value = _normalize_label_value(field, payload.get("value"), filename)
        manifest_key = f"packshots/{slug}/{MANIFEST}"
        try:
            labels = _manifest(storage, manifest_key)
        except Exception:
            labels = []
        label = next((row for row in labels if row.get("file") == filename), None)
        if label is None:
            label = {"file": filename, "annotation_name": os.path.splitext(filename)[0], "name": "", "id": ""}
            labels.append(label)
        duplicate = field == "annotation_name" and any(
            row is not label
            and (row.get("annotation_name") or os.path.splitext(row.get("file", ""))[0]).casefold() == value.casefold()
            for row in labels
        )
        label[field] = value
        label["updated_by"] = (getattr(request.state, "user", None) or {}).get("email", "")
        label["updated_at"] = datetime.now(timezone.utc).isoformat()
        label_etag = _label_etag(label, filename)
        record.update({field: value, "updated_by": label["updated_by"], "updated_at": label["updated_at"], "label_etag": label_etag})
        storage.put_json(manifest_key, labels)
        storage.put_json(_gallery_index_key("packshots", slug), gallery)
        _gallery_data_cache[(_storage_cache_key(storage), "packshots", slug)] = (time.monotonic() + 15, gallery)
        result = dict(record)
    if field == "annotation_name":
        result["duplicate_annotation_name"] = duplicate
    response = JSONResponse(result)
    return _response_timing(response, started, storage_calls=2)


@router.post("/packshots/{brand_slug}/upload")
async def upload_packshots(
    brand_slug: str,
    request: Request,
    files: list[UploadFile] = File(...),
    mapping_rows: str = Form(...),
):
    storage = request.app.state.storage
    _brand(storage, brand_slug)
    if not 1 <= len(files) <= MAX_FILES_PER_BATCH:
        raise HTTPException(400, f"Upload between 1 and {MAX_FILES_PER_BATCH} files per batch")
    try:
        rows = json.loads(mapping_rows)
    except (TypeError, json.JSONDecodeError):
        raise HTTPException(400, "Invalid mapping metadata")
    if not isinstance(rows, list) or len(rows) != len(files):
        raise HTTPException(400, "Mapping metadata does not match uploaded files")
    prefix = f"packshots/{brand_slug}"
    manifest_key = f"{prefix}/{MANIFEST}"
    manifest = _manifest(storage, manifest_key)
    by_file = {row.get("file"): row for row in manifest if isinstance(row, dict) and isinstance(row.get("file"), str)}
    added = ignored = failed = total_bytes = 0
    failures = []
    reserved = set()
    existing = set(storage.list(prefix))
    pending = []
    dimension_map = {}
    derivative_decisions = {}
    for upload, row in zip(files, rows):
        try:
            safe_name = _filename(upload.filename or "")
            if safe_name is None:
                ignored += 1
                continue
            data = await _read_image(upload)
            if data is None:
                ignored += 1
                continue
            total_bytes += len(data)
            if total_bytes > MAX_BATCH_BYTES:
                raise HTTPException(413, "Upload batch exceeds the 100 MB limit")
            collision = f"{prefix}/{safe_name}" in existing
            filename = _unique_key(storage, prefix, safe_name, reserved, existing)
            try:
                width, height = _image_dimensions(data)
                dimension_map[f"{prefix}/{filename}"] = {"width": width, "height": height}
                derivative_decisions[f"{prefix}/{filename}"] = "generated" if _derivative_gate(len(data), width, height) else "none"
            except Exception:
                pass
            if collision or not isinstance(row, dict):
                base = os.path.splitext(filename)[0]
                label = {"annotation_name": base, "name": base, "id": base}
            else:
                label = row
                label = {"annotation_name": str(label.get("annotation_name") or os.path.splitext(filename)[0]), "name": str(label.get("name") or ""), "id": str(label.get("id") or "")}
            pending.append((f"{prefix}/{filename}", data, "image/" + ("jpeg" if filename.lower().endswith((".jpg", ".jpeg")) else filename.rsplit(".", 1)[-1]), filename, label))
        except HTTPException as exc:
            failed += 1
            failures.append({"file": upload.filename or "unknown", "reason": exc.detail})
        except Exception as exc:
            failed += 1
            failures.append({"file": upload.filename or "unknown", "reason": str(exc)})
    write_failures = _store_objects(storage, [item[:4] for item in pending])
    for key, data, content_type, filename, label in pending:
        key = f"{prefix}/{filename}"
        if key in write_failures:
            failed += 1
            failures.append(write_failures[key])
        else:
            key = f"{prefix}/{filename}"
            by_file[filename] = {"file": filename, **label, "derivative": derivative_decisions.get(key), **dimension_map.get(key, {})}
            added += 1
    storage.put_json(manifest_key, list(by_file.values()))
    listed_entries = storage.list_entries(prefix)
    entries = {entry["key"]: entry for entry in listed_entries}
    for key, data, content_type, filename, label in pending:
        if key in write_failures or key not in entries:
            continue
        _schedule_derivatives(storage, "packshots", brand_slug, key, _version_for(entries[key]), data)
    _refresh_gallery(storage, "packshots", brand_slug, listed_entries, dimension_map, derivative_decisions)
    return {"added": added, "ignored": ignored, "failed": failed, "failures": failures}


@router.post("/packshots/{brand_slug}/{filename}/replace")
async def replace_packshot(brand_slug: str, filename: str, request: Request, image: UploadFile = File(...)):
    storage = request.app.state.storage
    _brand(storage, brand_slug)
    safe_name = _filename(filename)
    if safe_name is None or safe_name != filename:
        raise HTTPException(400, "Invalid packshot file name")
    prefix = f"packshots/{brand_slug}"
    key = f"{prefix}/{safe_name}"
    if not storage.exists(key):
        raise HTTPException(404, "Packshot not found")
    if os.path.splitext(image.filename or "")[1].lower() != os.path.splitext(safe_name)[1].lower():
        raise HTTPException(400, "Replacement image must use the same file extension")
    content = await _read_image(image)
    if content is None:
        raise HTTPException(400, "Replacement must be a supported image")
    storage.put(key, content, _image_content_type(key), cache_control="private, max-age=86400")
    width, height = _image_dimensions(content)
    decision = "generated" if _derivative_gate(len(content), width, height) else "none"
    manifest_key = f"{prefix}/{MANIFEST}"
    labels = _manifest(storage, manifest_key)
    for row in labels:
        if row.get("file") == safe_name:
            row.update({"derivative": decision, "width": width, "height": height})
            break
    storage.put_json(manifest_key, labels)
    listed_entries = storage.list_entries(prefix)
    _refresh_gallery(storage, "packshots", brand_slug, listed_entries, {key: {"width": width, "height": height}}, {key: decision})
    entries = _gallery_data_cache.get((_storage_cache_key(storage), "packshots", brand_slug), (0, {"images": []}))[1]["images"]
    entry = next((item for item in entries if item["key"] == key), None)
    if entry:
        _schedule_derivatives(storage, "packshots", brand_slug, key, _version_for(entry), content)
    return {"replaced": safe_name}


@router.get("/packshots/{brand_slug}/download")
def download_packshots(brand_slug: str, request: Request, selected: str = ""):
    storage = request.app.state.storage
    brand = _brand(storage, brand_slug)
    prefix = f"packshots/{brand_slug}"
    manifest = _manifest(storage, f"{prefix}/{MANIFEST}")
    rows = {row.get("file"): row for row in manifest if isinstance(row, dict)}
    requested = [name for name in selected.split(",") if name] if selected else None
    available = [key.rsplit("/", 1)[-1] for key in storage.list(prefix) if key.rsplit("/", 1)[-1] != MANIFEST]
    if requested is not None:
        if len(requested) != len(set(requested)) or any(name not in available for name in requested):
            raise HTTPException(400, "Invalid selected packshot list")
        available = requested
    if not available:
        raise HTTPException(404, "No packshots to download")
    temp = tempfile.NamedTemporaryFile(prefix="packshots-", suffix=".zip", delete=False)
    temp_path = temp.name
    used = set()
    try:
        with temp, zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for filename in available:
                row = rows.get(filename, {})
                label = row.get("annotation_name") or os.path.splitext(filename)[0]
                safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", str(label)).strip("._-") or os.path.splitext(filename)[0]
                zip_name = safe_stem + os.path.splitext(filename)[1]
                stem, ext = os.path.splitext(zip_name)
                candidate = zip_name
                suffix = 1
                while candidate.casefold() in used:
                    candidate = f"{stem}_{suffix}{ext}"
                    suffix += 1
                used.add(candidate.casefold())
                with storage.open_for_zip(f"{prefix}/{filename}") as source:
                    with archive.open(candidate, "w") as destination:
                        while chunk := source.read(1024 * 1024):
                            destination.write(chunk)
    except Exception:
        try:
            os.unlink(temp_path)
        except OSError:
            pass
        raise

    def stream_zip():
        try:
            with open(temp_path, "rb") as source:
                while chunk := source.read(1024 * 1024):
                    yield chunk
        finally:
            try:
                os.unlink(temp_path)
            except OSError:
                pass

    download_name = f"{brand.get('slug', brand_slug)}_packshots.zip"
    return StreamingResponse(stream_zip(), media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="{download_name}"'})
